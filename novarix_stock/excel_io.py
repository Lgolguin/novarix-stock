from __future__ import annotations

import os
import unicodedata
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException

from novarix_stock.database import PlanLimitError, StockDatabase, StockError
from novarix_stock.images import ProductImageStore
from novarix_stock.models import ImportedProductDraft, Product

EXCEL_HEADERS = (
    "Foto del producto",
    "Nombre",
    "Precio de compra",
    "Precio de venta",
    "Stock total",
    "Stock actual",
    "Vendidos",
    "SKU",
    "Stock mínimo",
)


class ExcelFileError(Exception):
    """Error de Excel que puede mostrarse directamente al usuario."""


@dataclass(frozen=True, slots=True)
class ExcelImportResult:
    imported_count: int
    duplicate_skus: tuple[str, ...]
    row_errors: tuple[str, ...]
    plan_limit_reached: bool = False


class StockExcelService:
    def __init__(
        self,
        database: StockDatabase,
        image_store: ProductImageStore,
    ) -> None:
        self.database = database
        self.image_store = image_store

    def import_file(self, source_path: str | Path) -> ExcelImportResult:
        source = Path(source_path)
        try:
            workbook = load_workbook(source, read_only=True, data_only=True)
        except (OSError, BadZipFile, InvalidFileException, ValueError) as error:
            raise ExcelFileError("El archivo no es un Excel .xlsx válido o no se puede abrir.") from error

        try:
            worksheet = workbook.active
            rows = worksheet.iter_rows(values_only=True)
            try:
                raw_headers = next(rows)
            except StopIteration as error:
                raise ExcelFileError("El archivo Excel está vacío.") from error
            header_indexes = self._map_headers(raw_headers)

            existing_skus = {product.sku.casefold() for product in self.database.list_products()}
            duplicate_skus: list[str] = []
            row_errors: list[str] = []
            imported_count = 0
            plan_limit_reached = False

            for row_number, row in enumerate(rows, start=2):
                if self._is_empty_row(row):
                    continue
                try:
                    draft = self._parse_row(row, row_number, header_indexes, source.parent)
                except ExcelFileError as error:
                    row_errors.append(str(error))
                    continue

                normalized_sku = draft.sku.casefold()
                if normalized_sku and normalized_sku in existing_skus:
                    duplicate_skus.append(draft.sku)
                    continue

                managed_photo = None
                try:
                    if draft.photo_path:
                        managed_photo = self.image_store.import_image(draft.photo_path)
                        draft = ImportedProductDraft(
                            sku=draft.sku,
                            name=draft.name,
                            purchase_price_cents=draft.purchase_price_cents,
                            sale_price_cents=draft.sale_price_cents,
                            stock_total=draft.stock_total,
                            stock_current=draft.stock_current,
                            sold=draft.sold,
                            minimum_stock=draft.minimum_stock,
                            photo_path=managed_photo,
                        )
                    imported_product = self.database.import_product(draft)
                except PlanLimitError:
                    self.image_store.delete_managed_image(managed_photo)
                    plan_limit_reached = True
                    break
                except (StockError, OSError) as error:
                    self.image_store.delete_managed_image(managed_photo)
                    row_errors.append(f"Fila {row_number}: {error}")
                    continue

                existing_skus.add(imported_product.sku.casefold())
                imported_count += 1

            return ExcelImportResult(
                imported_count=imported_count,
                duplicate_skus=tuple(duplicate_skus),
                row_errors=tuple(row_errors),
                plan_limit_reached=plan_limit_reached,
            )
        finally:
            workbook.close()

    def export_file(self, destination_path: str | Path) -> None:
        destination = Path(destination_path)
        if destination.suffix.lower() != ".xlsx":
            destination = destination.with_suffix(".xlsx")
        destination.parent.mkdir(parents=True, exist_ok=True)

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Stock"
        worksheet.sheet_view.showGridLines = False
        worksheet.freeze_panes = "A2"
        worksheet.append(EXCEL_HEADERS)

        for product in self.database.list_products():
            worksheet.append(self._export_row(product))

        self._format_worksheet(worksheet)
        temporary = destination.with_name(f".{destination.stem}.tmp.xlsx")
        try:
            workbook.save(temporary)
            os.replace(temporary, destination)
        except (OSError, ValueError) as error:
            try:
                temporary.unlink(missing_ok=True)
            except OSError:
                pass
            raise ExcelFileError(
                "No se pudo guardar el Excel. Verificá que el archivo no esté abierto."
            ) from error
        finally:
            workbook.close()

    @staticmethod
    def _map_headers(raw_headers: tuple[object, ...]) -> dict[str, int]:
        available: dict[str, int] = {}
        for index, value in enumerate(raw_headers):
            if value is not None:
                available[StockExcelService._normalize_header(str(value))] = index

        missing = [
            header for header in EXCEL_HEADERS
            if StockExcelService._normalize_header(header) not in available
        ]
        if missing:
            raise ExcelFileError(
                "Faltan columnas obligatorias: " + ", ".join(missing) + "."
            )
        return {
            header: available[StockExcelService._normalize_header(header)]
            for header in EXCEL_HEADERS
        }

    @staticmethod
    def _normalize_header(value: str) -> str:
        normalized = unicodedata.normalize("NFKD", value.strip().casefold())
        return " ".join(
            "".join(character for character in normalized if not unicodedata.combining(character)).split()
        )

    @staticmethod
    def _is_empty_row(row: tuple[object, ...]) -> bool:
        return all(value is None or (isinstance(value, str) and not value.strip()) for value in row)

    @classmethod
    def _parse_row(
        cls,
        row: tuple[object, ...],
        row_number: int,
        indexes: dict[str, int],
        workbook_dir: Path,
    ) -> ImportedProductDraft:
        def value(header: str) -> object:
            index = indexes[header]
            return row[index] if index < len(row) else None

        name = cls._required_text(value("Nombre"), "Nombre", row_number)
        sku = cls._optional_text(value("SKU"))
        purchase_cents = cls._price_to_cents(value("Precio de compra"), "Precio de compra", row_number)
        sale_cents = cls._price_to_cents(value("Precio de venta"), "Precio de venta", row_number)
        stock_total = cls._non_negative_integer(value("Stock total"), "Stock total", row_number)
        stock_current = cls._non_negative_integer(value("Stock actual"), "Stock actual", row_number)
        sold = cls._non_negative_integer(value("Vendidos"), "Vendidos", row_number)
        minimum = cls._non_negative_integer(value("Stock mínimo"), "Stock mínimo", row_number)

        if stock_total != stock_current + sold:
            raise ExcelFileError(
                f"Fila {row_number}: Stock total debe ser igual a Stock actual más Vendidos."
            )

        photo_path = None
        raw_photo = value("Foto del producto")
        if raw_photo is not None and str(raw_photo).strip():
            photo_candidate = Path(str(raw_photo).strip())
            if not photo_candidate.is_absolute():
                photo_candidate = workbook_dir / photo_candidate
            if not photo_candidate.is_file():
                raise ExcelFileError(f"Fila {row_number}: la fotografía indicada no existe.")
            photo_path = str(photo_candidate.resolve())

        return ImportedProductDraft(
            sku=sku,
            name=name,
            purchase_price_cents=purchase_cents,
            sale_price_cents=sale_cents,
            stock_total=stock_total,
            stock_current=stock_current,
            sold=sold,
            minimum_stock=minimum,
            photo_path=photo_path,
        )

    @staticmethod
    def _required_text(value: object, field: str, row_number: int) -> str:
        if value is None:
            raise ExcelFileError(f"Fila {row_number}: {field} es obligatorio.")
        if isinstance(value, float) and value.is_integer():
            text = str(int(value))
        else:
            text = str(value).strip()
        if not text:
            raise ExcelFileError(f"Fila {row_number}: {field} es obligatorio.")
        return text

    @staticmethod
    def _optional_text(value: object) -> str:
        if value is None:
            return ""
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        return str(value).strip()

    @classmethod
    def _price_to_cents(cls, value: object, field: str, row_number: int) -> int:
        decimal = cls._decimal_value(value, field, row_number)
        if decimal < 0:
            raise ExcelFileError(f"Fila {row_number}: {field} no puede ser negativo.")
        return int((decimal * 100).quantize(Decimal("1"), rounding=ROUND_HALF_UP))

    @classmethod
    def _non_negative_integer(cls, value: object, field: str, row_number: int) -> int:
        decimal = cls._decimal_value(value, field, row_number)
        if decimal < 0 or decimal != decimal.to_integral_value():
            raise ExcelFileError(
                f"Fila {row_number}: {field} debe ser un número entero no negativo."
            )
        return int(decimal)

    @staticmethod
    def _decimal_value(value: object, field: str, row_number: int) -> Decimal:
        if value is None or isinstance(value, bool):
            raise ExcelFileError(f"Fila {row_number}: {field} debe ser numérico.")
        text = str(value).strip().replace("$", "").replace(" ", "")
        if not text:
            raise ExcelFileError(f"Fila {row_number}: {field} debe ser numérico.")
        if "," in text and "." in text:
            if text.rfind(",") > text.rfind("."):
                text = text.replace(".", "").replace(",", ".")
            else:
                text = text.replace(",", "")
        elif "," in text:
            text = text.replace(",", ".")
        try:
            result = Decimal(text)
        except InvalidOperation as error:
            raise ExcelFileError(f"Fila {row_number}: {field} debe ser numérico.") from error
        if not result.is_finite():
            raise ExcelFileError(f"Fila {row_number}: {field} debe ser numérico.")
        return result

    @staticmethod
    def _export_row(product: Product) -> tuple[object, ...]:
        return (
            product.photo_path or "",
            product.name,
            product.purchase_price_cents / 100,
            product.sale_price_cents / 100,
            product.stock_total,
            product.stock_current,
            product.sold,
            product.sku,
            product.minimum_stock,
        )

    @staticmethod
    def _format_worksheet(worksheet: object) -> None:
        header_fill = PatternFill("solid", fgColor="18243A")
        header_font = Font(name="Segoe UI", bold=True, color="4EE5FF")
        body_font = Font(name="Segoe UI", color="172033")
        thin_border = Border(bottom=Side(style="thin", color="D8E0EC"))

        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        worksheet.row_dimensions[1].height = 28

        for row in worksheet.iter_rows(min_row=2):
            for cell in row:
                cell.font = body_font
                cell.border = thin_border
                cell.alignment = Alignment(vertical="center")
        worksheet.auto_filter.ref = worksheet.dimensions

        widths = (34, 28, 18, 18, 14, 14, 12, 18, 14)
        for column, width in enumerate(widths, start=1):
            worksheet.column_dimensions[get_column_letter(column)].width = width

        for column in (3, 4):
            for cell in worksheet[get_column_letter(column)][1:]:
                cell.number_format = '"$"#,##0.00'
        for column in (5, 6, 7, 9):
            for cell in worksheet[get_column_letter(column)][1:]:
                cell.number_format = "#,##0"
