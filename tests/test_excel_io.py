from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from novarix_stock.database import StockDatabase
from novarix_stock.excel_io import EXCEL_HEADERS, ExcelFileError, StockExcelService
from novarix_stock.images import ProductImageStore
from novarix_stock.models import ProductDraft


class ExcelIoTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)
        self.db_path = self.root / "stock.db"
        self.database = StockDatabase(self.db_path)
        self.image_store = ProductImageStore(self.root / "images")
        self.service = StockExcelService(self.database, self.image_store)

    def tearDown(self) -> None:
        self.temp_dir.cleanup()

    def _write_workbook(self, path: Path, rows: list[tuple[object, ...]]) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(EXCEL_HEADERS)
        for row in rows:
            sheet.append(row)
        workbook.save(path)
        workbook.close()

    def test_import_valid_skips_duplicate_and_invalid_row(self) -> None:
        self.database.add_product(
            ProductDraft(
                sku="EXIST-1",
                name="Existente",
                purchase_price_cents=1000,
                sale_price_cents=1500,
                minimum_stock=1,
                initial_stock=2,
            )
        )
        source = self.root / "importacion.xlsx"
        self._write_workbook(
            source,
            [
                ("", "Producto válido", 12.50, 20, 12, 8, 4, "NOV-100", 3),
                ("", "Duplicado", 10, 15, 2, 2, 0, "exist-1", 1),
                ("", "Fila inválida", 10, 15, 9, 8, 4, "BAD-1", 2),
            ],
        )

        result = self.service.import_file(source)

        self.assertEqual(result.imported_count, 1)
        self.assertEqual(result.duplicate_skus, ("exist-1",))
        self.assertEqual(len(result.row_errors), 1)
        self.assertIn("Fila 4", result.row_errors[0])
        imported = {product.sku: product for product in self.database.list_products()}
        self.assertEqual(set(imported), {"EXIST-1", "NOV-100"})
        self.assertEqual(
            (
                imported["NOV-100"].purchase_price_cents,
                imported["NOV-100"].stock_total,
                imported["NOV-100"].stock_current,
                imported["NOV-100"].sold,
            ),
            (1250, 12, 8, 4),
        )

        reopened = StockDatabase(self.db_path)
        self.assertEqual(len(reopened.list_products()), 2)

    def test_import_generates_distinct_skus_for_empty_cells(self) -> None:
        source = self.root / "sku_vacios.xlsx"
        self._write_workbook(
            source,
            [
                ("", "Sin SKU 1", 10, 15, 2, 2, 0, "", 0),
                ("", "Sin SKU 2", 10, 15, 3, 3, 0, None, 0),
                ("", "SKU manual", 10, 15, 1, 1, 0, "MANUAL-XLSX", 0),
            ],
        )

        result = self.service.import_file(source)

        self.assertEqual(result.imported_count, 3)
        self.assertEqual(result.duplicate_skus, ())
        self.assertEqual(result.row_errors, ())
        products = {product.name: product.sku for product in self.database.list_products()}
        self.assertEqual(products["Sin SKU 1"], "STK-0001")
        self.assertEqual(products["Sin SKU 2"], "STK-0002")
        self.assertEqual(products["SKU manual"], "MANUAL-XLSX")

    def test_missing_column_aborts_without_importing(self) -> None:
        source = self.root / "faltan_columnas.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(EXCEL_HEADERS[:-1])
        sheet.append(("", "Producto", 10, 15, 2, 2, 0, "NOV-1"))
        workbook.save(source)
        workbook.close()

        with self.assertRaisesRegex(ExcelFileError, "Stock mínimo"):
            self.service.import_file(source)
        self.assertEqual(self.database.list_products(), [])

    def test_export_is_readable_and_keeps_numeric_cells(self) -> None:
        self.database.add_product(
            ProductDraft(
                sku="EXP-001",
                name="Producto exportado",
                purchase_price_cents=12345,
                sale_price_cents=19999,
                minimum_stock=2,
                initial_stock=7,
            )
        )
        destination = self.root / "stock_exportado.xlsx"

        self.service.export_file(destination)

        workbook = load_workbook(destination, data_only=False)
        sheet = workbook["Stock"]
        self.assertEqual(tuple(cell.value for cell in sheet[1]), EXCEL_HEADERS)
        self.assertEqual(sheet["B2"].value, "Producto exportado")
        self.assertEqual(sheet["C2"].value, 123.45)
        self.assertEqual(sheet["E2"].value, 7)
        self.assertEqual(sheet["H2"].value, "EXP-001")
        self.assertEqual(sheet["C2"].number_format, '"$"#,##0.00')
        self.assertGreater(sheet.column_dimensions["B"].width, 20)
        self.assertEqual(sheet.freeze_panes, "A2")
        workbook.close()

    def test_free_import_stops_at_limit_and_resumes_after_deletion(self) -> None:
        for index in range(9):
            self.database.add_product(
                ProductDraft(
                    sku=f"BASE-{index:02d}",
                    name=f"Base {index}",
                    purchase_price_cents=1000,
                    sale_price_cents=1500,
                    minimum_stock=0,
                )
            )
        source = self.root / "limite_free.xlsx"
        self._write_workbook(
            source,
            [
                ("", "Importado 1", 10, 15, 0, 0, 0, "IMP-1", 0),
                ("", "Importado 2", 10, 15, 0, 0, 0, "IMP-2", 0),
                ("", "Importado 3", 10, 15, 0, 0, 0, "IMP-3", 0),
            ],
        )

        result = self.service.import_file(source)

        self.assertEqual(result.imported_count, 1)
        self.assertTrue(result.plan_limit_reached)
        self.assertEqual(len(self.database.list_products()), 10)
        self.assertIsNotNone(next((p for p in self.database.list_products() if p.sku == "IMP-1"), None))
        self.assertIsNone(next((p for p in self.database.list_products() if p.sku == "IMP-2"), None))

        product_to_delete = next(p for p in self.database.list_products() if p.sku == "BASE-00")
        self.database.delete_product(product_to_delete.id)
        resumed_source = self.root / "reanudar_free.xlsx"
        self._write_workbook(
            resumed_source,
            [("", "Importado nuevo", 10, 15, 0, 0, 0, "IMP-NEW", 0)],
        )

        resumed = self.service.import_file(resumed_source)

        self.assertEqual(resumed.imported_count, 1)
        self.assertFalse(resumed.plan_limit_reached)
        self.assertEqual(len(self.database.list_products()), 10)


if __name__ == "__main__":
    unittest.main()
